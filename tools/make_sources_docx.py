"""Собирает «Список использованных источников.docx» из data/sources.json.

Библиографические записи — по ГОСТ Р 7.0.100–2018, оформление страницы — по ГОСТ 7.32–2017:
Times New Roman 14 пт, полуторный интервал, абзацный отступ 1,25 см, поля 30/15/20/20 мм,
нумерация страниц внизу по центру. Источники на кириллице идут по алфавиту, затем на латинице.

    .venv\\Scripts\\python tools\\make_sources_docx.py
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Список использованных источников.docx"
FONT = "Times New Roman"


def sort_key(entry: str) -> tuple[int, str]:
    text = entry.lower().replace("ё", "е")
    text = re.sub(r"^[\[«\"'(]+", "", text)
    is_latin = bool(re.match(r"[a-z]", text))
    return (1 if is_latin else 0, text)


def set_fonts(rpr_parent) -> None:
    """Times New Roman для всех письменностей, включая кириллицу."""
    rpr = rpr_parent.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def add_numbering(document) -> int:
    """Нумерация «1.» с абзацным отступом 1,25 см и переносом строк к левому полю."""
    numbering = document.part.numbering_part.element
    used = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    used += [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    new_id = str(max(used, default=0) + 1)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), new_id)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:suff", "space"),
                       ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), value)
        lvl.append(el)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:firstLine"), str(round(Cm(1.25).twips)))
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), new_id)
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), new_id)
    num.append(ref)
    # порядок по схеме: все w:abstractNum, затем все w:num
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(abstract)
        nums[-1].addnext(num)
    else:
        numbering.append(abstract)
        numbering.append(num)
    return int(new_id)


def number_paragraph(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(nid)
    ppr.insert(0, num_pr)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text.strip() == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def main() -> None:
    sources = json.loads((ROOT / "data" / "sources.json").read_text(encoding="utf-8"))
    entries = sorted((s["gost"] for s in sources), key=sort_key)

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(30), Mm(15)
    section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    section.footer_distance = Mm(10)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    set_fonts(normal.element)
    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Cm(1.25)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.widow_control = True

    document.core_properties.title = "Список использованных источников"
    document.core_properties.subject = "Голосовой путеводитель по Карелии"
    document.core_properties.language = "ru-RU"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(21)  # одна пустая строка при полуторном интервале
    title.paragraph_format.keep_with_next = True
    run = title.add_run("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    run.bold = True

    num_id = add_numbering(document)
    for entry in entries:
        # неразрывный пробел перед « : », « – », « // », « / » — строка не начнётся со знака
        for mark in (":", "–", "//", "/"):
            entry = entry.replace(f" {mark} ", f" {mark} ")
        paragraph = document.add_paragraph(entry)
        number_paragraph(paragraph, num_id)

    add_page_number(section)
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:  # в шаблоне python-docx у масштаба нет обязательного атрибута
        zoom.set(qn("w:percent"), "100")
    document.save(OUT)
    print(f"Сохранено: {OUT} ({len(entries)} источников)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
